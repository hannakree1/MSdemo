"""PRG (rühmagrantide) ülevaade otse ETIS-i avalikust API-st.

Käivitamine:  python3 prg_ulevaade.py
Tulemus:      valjund/PRG_ulevaade_<kuupäev>.xlsx

Metoodika:
- Andmed: https://www.etis.ee:2346/api/project/getitems, filter ProgrammeCode=PRG.
- Algusaasta: projekt arvestatakse aastasse, mil see algas; summa = projekti kogurahastus.
- Eelarveaasta: iga rahastusperioodi summa jagatakse aastate vahel päevade järgi;
  aktiivne projekt = projekt, millel on sel aastal rahastust.
- Asutused: koondatud peaasutuse tasemele (nimi enne esimest koma). Mitme asutusega
  projekt jagatakse asutuste vahel võrdselt.
- Valdkonnad: Frascati klassifikaator. Projekt (ja selle rahastus) jagatakse valdkondade
  vahel osakaalude järgi, seega murdosad liituvad kokku täpselt projektide arvuks.
  Frascati valdkondadel ETIS-is osakaalu pole; ETIS hoiab need paarides ETIS-i
  valdkondadega, seega võetakse osakaal samal kohal olevalt ETIS-i valdkonnalt.
"""

import json
import urllib.request
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

API = "https://www.etis.ee:2346/api/project/getitems?Format=json&ProgrammeCode=PRG&Take={take}&Skip={skip}"
LEHT = 200

# ETIS-i kirjed, mis ei ole iseseisvad PRG projektid: Guid -> põhjus
VALJA_JAETUD = {
    "8fef4bfa-2ba2-456e-ac5a-1ab12772354f": "Numbrita Tallinna Ülikooli partnerikirje projektile PRG1919 (topeltloendus)",
    "ebecf66e-1ec0-4923-a01b-cacf5c4abb92": "ETIS-i süsteemi testkirje",
}

POHIVALDKONNAD = {
    "1": "Loodusteadused",
    "2": "Tehnika ja tehnoloogia",
    "3": "Arsti- ja terviseteadused",
    "4": "Põllumajandusteadused",
    "5": "Sotsiaalteadused",
    "6": "Humanitaarteadused",
}


def lae_projektid():
    projektid, skip = [], 0
    while True:
        with urllib.request.urlopen(API.format(take=LEHT, skip=skip), timeout=120) as vastus:
            osa = json.load(vastus)
        projektid += osa
        if len(osa) < LEHT:
            return projektid
        skip += LEHT


def kp(tekst):
    p, k, a = tekst.split(".")
    return date(int(a), int(k), int(p))


def summad_aastate_kaupa(projekt):
    """Rahastus aastate kaupa; aastavahetust ületav periood jagatakse päevade järgi."""
    aastad = defaultdict(float)
    for periood in projekt["FinancingPeriods"]:
        algus, lopp, summa = kp(periood["PeriodStartDate"]), kp(periood["PeriodEndDate"]), periood["Amount"] or 0
        paevi = (lopp - algus).days + 1
        for aasta in range(algus.year, lopp.year + 1):
            a = max(algus, date(aasta, 1, 1))
            l = min(lopp, date(aasta, 12, 31))
            aastad[aasta] += summa * ((l - a).days + 1) / paevi
    return aastad


def valdkonnad(projekt):
    """Frascati valdkonnad koos osakaaluga, mis võetakse paarilise ETIS-i valdkonna juurest.
    Sama Frascati kood mitme ETIS-i valdkonna all liidetakse kokku."""
    osad, nimed = defaultdict(float), {}
    for e, f in zip(projekt["ResearchAreasEtis"], projekt["ResearchAreasFrascati"], strict=True):
        osad[f["Code"]] += (e["Percent"] or 0) / 100
        nimed[f["Code"]] = f["Name"]
    return [(kood, nimed[kood], osa) for kood, osa in osad.items()]


def asutused(projekt):
    return sorted({i["Name"].split(",")[0].strip() for i in projekt["Institutions"]})


def vastutavad(projekt):
    return ", ".join(x["Name"] for x in projekt["Persons"] if x["RoleName"] == "Vastutav täitja")


def arvuta(projektid):
    algusaasta = defaultdict(lambda: [0, 0.0])
    eelarveaasta = defaultdict(lambda: [0, 0.0])
    asutus = defaultdict(lambda: [0.0, 0.0])
    suur = defaultdict(lambda: [0.0, 0.0])
    ala = defaultdict(lambda: [0.0, 0.0])
    ala_nimi = {}

    for p in projektid:
        kokku = p["FinancingInPeriodsTotal"] or 0

        a = algusaasta[kp(p["ProjectStartDate"]).year]
        a[0] += 1
        a[1] += kokku

        for aasta, summa in summad_aastate_kaupa(p).items():
            if summa > 0:
                eelarveaasta[aasta][0] += 1
                eelarveaasta[aasta][1] += summa

        asut = asutused(p)
        for nimi in asut:
            asutus[nimi][0] += 1 / len(asut)
            asutus[nimi][1] += kokku / len(asut)

        for kood, nimi, osa in valdkonnad(p):
            ala_nimi[kood] = nimi
            for rida in (ala[kood], suur[kood.split(".")[0]]):
                rida[0] += osa
                rida[1] += kokku * osa

    return algusaasta, eelarveaasta, asutus, suur, ala, ala_nimi


PAIS = Font(bold=True, color="FFFFFF")
PAIS_TAUST = PatternFill("solid", fgColor="2F5597")
KOKKU = Font(bold=True)
EUR = '#,##0 "€"'
ARV = "0.00"
PROTS = "0.0%"


def leht(wb, nimi, pais, read, vormingud, laiused, kokku_rida=None):
    ws = wb.create_sheet(nimi)
    ws.append(pais)
    for lahter in ws[1]:
        lahter.font, lahter.fill = PAIS, PAIS_TAUST
        lahter.alignment = Alignment(wrap_text=True, vertical="center")
    for rida in read:
        ws.append(rida)
    if kokku_rida:
        ws.append(kokku_rida)
        for lahter in ws[ws.max_row]:
            lahter.font = KOKKU
    for veerg, vorming in enumerate(vormingud, start=1):
        if vorming:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, veerg).number_format = vorming
    for veerg, laius in enumerate(laiused, start=1):
        ws.column_dimensions[get_column_letter(veerg)].width = laius
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(pais))}{len(read) + 1}"
    return ws


def osakaaluga(andmed, koguarv, kogusumma):
    return [(v, arv, arv / koguarv, summa, summa / kogusumma) for v, (arv, summa) in andmed]


def kirjuta_excel(koik, tee):
    projektid = [p for p in koik if p["Guid"] not in VALJA_JAETUD]
    valja = [p for p in koik if p["Guid"] in VALJA_JAETUD]
    algusaasta, eelarveaasta, asutus, suur, ala, ala_nimi = arvuta(projektid)
    n = len(projektid)
    raha = sum(p["FinancingInPeriodsTotal"] or 0 for p in projektid)
    numbrita = [p for p in projektid if not p["FinancierProjectNr"]]

    wb = Workbook()
    ws = wb.active
    ws.title = "Kokkuvõte"
    ws.append(["PRG ülevaade (ETIS)"])
    ws["A1"].font = Font(bold=True, size=14)
    for rida in [
        [],
        ["Andmed laaditud", date.today().strftime("%d.%m.%Y")],
        ["Allikas", "ETIS avalik API, ProgrammeCode=PRG"],
        ["PRG projekte kokku", n],
        ["Rahastus kokku", raha],
        ["Projektinumbrita kirjeid", len(numbrita)],
        [],
        ["Metoodika"],
        ["Algusaasta", "Projekt loetakse aastasse, mil see algas; summa on projekti kogurahastus."],
        ["Eelarveaasta", "Rahastus jagatud aastate vahel päevade järgi; aktiivne = sel aastal rahastust."],
        ["Asutused", "Peaasutuse tasemel; mitme asutusega projekt jagatud asutuste vahel võrdselt."],
        ["Valdkonnad", "Frascati klassifikaator; projekt ja rahastus jagatud valdkondade vahel osakaalude järgi."],
        ["", "Osakaal võetud paariliselt ETIS-i valdkonnalt, sest Frascati valdkondadel ETIS-is osakaalu pole."],
        ["Märkus", "Jooksva ja tulevaste aastate summad võivad olla plaanitud, mitte välja makstud."],
        [],
        ["Välja jäetud kirjed"],
    ]:
        ws.append(rida)
    ws.cell(ws.max_row, 1).font = KOKKU
    for p in valja:
        summa = f'{p["FinancingInPeriodsTotal"]:,.0f}'.replace(",", " ")
        ws.append([p["FinancierProjectNr"] or "(number puudub)",
                   f'{p["Title"]} ({summa} €) – {VALJA_JAETUD[p["Guid"]]}'])
    ws["B6"].number_format = EUR
    ws["A9"].font = KOKKU
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90

    read = [(a, arv, summa) for a, (arv, summa) in sorted(algusaasta.items())]
    leht(wb, "Algusaasta", ["Algusaasta", "Alustatud PRG-sid", "Nende kogurahastus"],
         read, [None, None, EUR], [12, 18, 20], ["Kokku", n, raha])

    read = [(a, arv, summa) for a, (arv, summa) in sorted(eelarveaasta.items())]
    leht(wb, "Eelarveaasta", ["Eelarveaasta", "Aktiivseid PRG-sid", "Rahastus aastal"],
         read, [None, None, EUR], [12, 18, 20],
         ["Kokku", None, sum(r[2] for r in read)])

    read = osakaaluga(sorted(asutus.items(), key=lambda x: -x[1][1]), n, raha)
    leht(wb, "Asutused", ["Asutus", "PRG-sid", "Osakaal arvust", "Rahastus", "Osakaal rahast"],
         read, [None, ARV, PROTS, EUR, PROTS], [60, 10, 14, 18, 14], ["Kokku", n, 1, raha, 1])

    read = [(f"{k}. {POHIVALDKONNAD[k]}", *r) for k, *r in
            osakaaluga(sorted(suur.items()), n, raha)]
    leht(wb, "Frascati põhivaldkonnad", ["Põhivaldkond", "PRG-sid (osakaalu järgi)", "Osakaal arvust", "Rahastus", "Osakaal rahast"],
         read, [None, ARV, PROTS, EUR, PROTS], [34, 14, 14, 18, 14], ["Kokku", n, 1, raha, 1])

    jarjestus = sorted(ala, key=lambda k: tuple(int(x) for x in k.split(".")))
    read = [(k, ala_nimi[k], POHIVALDKONNAD[k.split(".")[0]], arv, arv / n, summa, summa / raha)
            for k in jarjestus for arv, summa in [ala[k]]]
    leht(wb, "Frascati alavaldkonnad", ["Kood", "Alavaldkond", "Põhivaldkond", "PRG-sid (osakaalu järgi)", "Osakaal arvust", "Rahastus", "Osakaal rahast"],
         read, [None, None, None, ARV, PROTS, EUR, PROTS], [8, 60, 26, 14, 14, 18, 14],
         ["Kokku", None, None, n, 1, raha, 1])

    read = [(p["FinancierProjectNr"] or "(number puudub)", p["Title"], p["ProjectStartDate"], p["ProjectEndDate"],
             "; ".join(asutused(p)), vastutavad(p), p["FinancingInPeriodsTotal"] or 0,
             "; ".join(f"{kood} {nimi} ({osa:.0%})" for kood, nimi, osa in valdkonnad(p)))
            for p in sorted(projektid, key=lambda p: (kp(p["ProjectStartDate"]), p["FinancierProjectNr"]))]
    leht(wb, "Projektid", ["Nr", "Pealkiri", "Algus", "Lõpp", "Asutus", "Vastutav täitja", "Rahastus", "Frascati valdkonnad"],
         read, [None, None, None, None, None, None, EUR, None], [14, 60, 11, 11, 30, 28, 16, 70])

    tee.parent.mkdir(exist_ok=True)
    wb.save(tee)


if __name__ == "__main__":
    projektid = lae_projektid()
    tee = Path(__file__).parent / "valjund" / f"PRG_ulevaade_{date.today():%Y-%m-%d}.xlsx"
    kirjuta_excel(projektid, tee)
    valja = sum(p["Guid"] in VALJA_JAETUD for p in projektid)
    print(f"{len(projektid) - valja} PRG projekti ({valja} kirjet välja jäetud) -> {tee}")
